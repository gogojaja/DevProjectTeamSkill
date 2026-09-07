#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""xlsx_cell_fix.py — T-05 Excel 批量修复器

按修复指令 JSON 批量修改 xlsx 单元格值，输出修复后 xlsx + 执行日志。
解决痛点 P3（修复脚本一次性编写）。

CLI（跨平台）：
  py -3.11 tools/xlsx_cell_fix.py --input <xlsx> --fixes <修复指令JSON> --output <xlsx>
  py -3.11 tools/xlsx_cell_fix.py --input req.xlsx --fixes fixes.json --output fixed.xlsx

修复指令 JSON 格式：
  {"version":"1.0","source_file":"...","fixes":[
    {"defect_id":"DEF-001","description":"...","sheet":"Sheet1",
     "cells":[{"ref":"D21","old_value":"显示","new_value":"计算后回显"}]},
    {"defect_id":"DEF-003","sheet":"Sheet2","insert_row":10,
     "cells":[{"ref":"A10","value":"新增行"}]}
  ]}

输出：
  修复后 xlsx + 执行日志 JSON（含每步操作结果）

依赖：openpyxl（未安装时提示安装命令）
"""
import os
import sys
import io
import json
import copy
import shutil
import argparse
from datetime import datetime

# Windows 控制台 UTF-8 输出
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ROOT = os.environ.get("PROJECT_ROOT", os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def _now():
    return datetime.now().strftime("%Y-%m-%dT%H:%M:%S")


def apply_fixes(input_path, fixes_path, output_path=None):
    """按修复指令批量修改 xlsx。

    Args:
        input_path: 原始 xlsx 路径
        fixes_path: 修复指令 JSON 路径
        output_path: 输出 xlsx 路径（None 则自动生成）

    Returns:
        dict: 执行日志
    """
    if not HAS_OPENPYXL:
        print("错误: openpyxl 未安装。请运行: pip install openpyxl")
        sys.exit(1)

    for p in [input_path, fixes_path]:
        if not os.path.exists(p):
            print("错误: 文件不存在: %s" % p)
            sys.exit(1)

    # 加载修复指令
    with io.open(fixes_path, "r", encoding="utf-8") as f:
        fix_data = json.load(f)

    fixes = fix_data.get("fixes", [])
    if not fixes:
        print("警告: 修复指令为空")
        return {"total": 0, "success": 0, "failed": 0}

    # 自动备份原始文件
    backup_dir = os.path.join(ROOT, ".backup")
    os.makedirs(backup_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(backup_dir, "xlsx_fix_backup_%s_%s" % (
        os.path.basename(input_path), ts))
    shutil.copy2(input_path, backup_path)

    # 加载 xlsx
    wb = openpyxl.load_workbook(input_path)

    if not output_path:
        base = os.path.splitext(os.path.basename(input_path))[0]
        output_path = os.path.join(ROOT, "docs", "reviews", "%s_fixed.xlsx" % base)

    log = {
        "version": fix_data.get("version", "1.0"),
        "source_file": input_path,
        "backup_file": backup_path,
        "fix_time": _now(),
        "total_fixes": len(fixes),
        "success": 0,
        "failed": 0,
        "details": [],
    }

    for fix in fixes:
        defect_id = fix.get("defect_id", "未知")
        sheet_name = fix.get("sheet", "")
        description = fix.get("description", "")
        insert_row = fix.get("insert_row", None)
        cells = fix.get("cells", [])

        fix_result = {
            "defect_id": defect_id,
            "description": description,
            "cell_ops": [],
        }

        if sheet_name not in wb.sheetnames:
            fix_result["status"] = "失败"
            fix_result["error"] = "工作表不存在: %s" % sheet_name
            log["failed"] += 1
            log["details"].append(fix_result)
            continue

        ws = wb[sheet_name]

        # 插入行（如果有 insert_row）
        if insert_row:
            ws.insert_rows(insert_row)
            fix_result["cell_ops"].append({
                "op": "insert_row",
                "row": insert_row,
                "status": "成功",
            })

        # 修改单元格
        all_ok = True
        for cell_spec in cells:
            ref = cell_spec.get("ref", "")
            new_value = cell_spec.get("new_value", cell_spec.get("value", ""))
            old_value = cell_spec.get("old_value", "")

            try:
                cell = ws[ref]
                actual_old = str(cell.value) if cell.value is not None else ""

                # 如果指定了 old_value，验证当前值
                if old_value and actual_old != old_value:
                    fix_result["cell_ops"].append({
                        "op": "modify",
                        "ref": ref,
                        "expected_old": old_value,
                        "actual_old": actual_old,
                        "status": "跳过（值不匹配）",
                    })
                    continue

                cell.value = new_value
                fix_result["cell_ops"].append({
                    "op": "modify",
                    "ref": ref,
                    "old_value": actual_old,
                    "new_value": new_value,
                    "status": "成功",
                })
            except Exception as e:
                fix_result["cell_ops"].append({
                    "op": "modify",
                    "ref": ref,
                    "status": "失败",
                    "error": str(e),
                })
                all_ok = False

        fix_result["status"] = "成功" if all_ok else "部分失败"
        log["success"] += 1 if all_ok else 0
        log["failed"] += 0 if all_ok else 1
        log["details"].append(fix_result)

    # 保存修复后 xlsx
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    wb.save(output_path)

    # 写入执行日志
    log_path = os.path.splitext(output_path)[0] + "_fix_log.json"
    with io.open(log_path, "w", encoding="utf-8") as f:
        json.dump(log, f, ensure_ascii=False, indent=2)

    print("xlsx 批量修复完成：")
    print("  输入: %s" % input_path)
    print("  备份: %s" % backup_path)
    print("  输出: %s" % output_path)
    print("  日志: %s" % log_path)
    print("  修复指令: %d 条 | 成功: %d | 失败: %d" % (
        log["total_fixes"], log["success"], log["failed"]))
    return log


def main():
    ap = argparse.ArgumentParser(description="T-05 Excel 批量修复器")
    ap.add_argument("--input", required=True, help="原始 xlsx 文件路径")
    ap.add_argument("--fixes", required=True, help="修复指令 JSON 文件路径")
    ap.add_argument("--output", default=None, help="输出 xlsx 路径")
    args = ap.parse_args()
    apply_fixes(args.input, args.fixes, args.output)


if __name__ == "__main__":
    main()
