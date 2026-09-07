#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""xlsx_dump.py — T-01 Excel 全量导出器

读取 xlsx 全工作表内容 → 输出 txt/json（含单元格坐标、合并单元格信息）。
解决痛点 P1（PowerShell 中文编码问题）。

CLI（跨平台）：
  py -3.11 tools/xlsx_dump.py --input <xlsx路径> --output <输出路径> --format txt|json
  py -3.11 tools/xlsx_dump.py --input requirements/test.xlsx --output dump.txt

依赖：openpyxl（未安装时提示安装命令）
"""
import os
import sys
import io
import json
import argparse

# Windows 控制台 UTF-8 输出
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ROOT = os.environ.get("PROJECT_ROOT", os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def dump_xlsx(input_path, output_path=None, fmt="txt"):
    """导出 xlsx 全工作表内容。

    Args:
        input_path: xlsx 文件路径
        output_path: 输出路径（None 则自动生成）
        fmt: "txt" 或 "json"

    Returns:
        str: 输出文件路径
    """
    if not HAS_OPENPYXL:
        print("错误: openpyxl 未安装。请运行: pip install openpyxl")
        sys.exit(1)

    if not os.path.exists(input_path):
        print("错误: 文件不存在: %s" % input_path)
        sys.exit(1)

    if not output_path:
        base = os.path.splitext(os.path.basename(input_path))[0]
        ext = ".json" if fmt == "json" else ".txt"
        output_path = os.path.join(ROOT, "docs", "reviews", base + "_dump" + ext)

    wb = openpyxl.load_workbook(input_path, data_only=True)
    sheets_data = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        cells = []
        # 获取合并单元格信息
        merged = []
        for m in ws.merged_cells.ranges:
            merged.append(str(m))

        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells.append({
                        "ref": cell.coordinate,
                        "row": cell.row,
                        "col": cell.column,
                        "value": str(cell.value),
                    })

        sheets_data[sheet_name] = {
            "cells": cells,
            "merged": merged,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
        }

    if fmt == "json":
        with io.open(output_path, "w", encoding="utf-8") as f:
            json.dump(sheets_data, f, ensure_ascii=False, indent=2)
    else:
        with io.open(output_path, "w", encoding="utf-8") as f:
            for sheet_name, data in sheets_data.items():
                f.write("=" * 60 + "\n")
                f.write("工作表: %s (%d 行 x %d 列)\n" % (sheet_name, data["max_row"], data["max_col"]))
                if data["merged"]:
                    f.write("合并单元格: %s\n" % ", ".join(data["merged"]))
                f.write("-" * 60 + "\n")
                for cell in data["cells"]:
                    f.write("  %s: %s\n" % (cell["ref"], cell["value"]))
                f.write("\n")

    total_cells = sum(len(d["cells"]) for d in sheets_data.values())
    print("xlsx 导出完成：")
    print("  输入: %s" % input_path)
    print("  输出: %s" % output_path)
    print("  工作表: %d 个, 单元格: %d 个" % (len(sheets_data), total_cells))
    return output_path


def main():
    ap = argparse.ArgumentParser(description="T-01 Excel 全量导出器")
    ap.add_argument("--input", required=True, help="xlsx 文件路径")
    ap.add_argument("--output", default=None, help="输出路径")
    ap.add_argument("--format", choices=["txt", "json"], default="txt", help="输出格式")
    args = ap.parse_args()
    dump_xlsx(args.input, args.output, args.format)


if __name__ == "__main__":
    main()
